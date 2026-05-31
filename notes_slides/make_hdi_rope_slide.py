"""Generate a single PowerPoint slide explaining the HDI+ROPE method."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import matplotlib.pyplot as plt
import numpy as np
import os

# ── Create the visual diagram ─────────────────────────────────────────────────
fig, axes = plt.subplots(1, 2, figsize=(8, 2.2))

for ax, case, color, hdi_pos, label in [
    (axes[0], "DE: higher in k₁", "#2ca02c", (0.6, 1.8), "HDI entirely right of ROPE"),
    (axes[1], "Not DE", "#d62728", (-0.4, 0.8), "HDI overlaps ROPE"),
]:
    ax.set_xlim(-2.5, 2.5)
    ax.set_ylim(-0.5, 1.5)
    ax.axvspan(-0.5, 0.5, alpha=0.15, color='grey')
    ax.axvline(-0.5, color='grey', ls='--', lw=1)
    ax.axvline(0.5, color='grey', ls='--', lw=1)
    ax.axvline(0, color='grey', ls=':', lw=0.8)
    ax.plot(hdi_pos, [0.7, 0.7], color=color, lw=6, solid_capstyle='round')
    ax.plot(hdi_pos[0], 0.7, '|', color=color, ms=15, mew=2)
    ax.plot(hdi_pos[1], 0.7, '|', color=color, ms=15, mew=2)
    ax.text(np.mean(hdi_pos), 0.95, "95% HDI", ha='center', fontsize=10, color=color, fontweight='bold')
    ax.text(0, -0.3, "ROPE [−ε, +ε]", ha='center', fontsize=9, color='grey')
    ax.set_title(case, fontsize=12, fontweight='bold', color=color)
    ax.text(np.mean(hdi_pos), 0.35, label, ha='center', fontsize=8.5, style='italic')
    ax.set_yticks([])
    ax.set_xticks([-0.5, 0, 0.5])
    ax.set_xticklabels(["−ε", "0", "+ε"], fontsize=9)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_visible(False)
    ax.set_xlabel("δ_g = B_{k₁,g} − B_{k₂,g}", fontsize=9)

plt.tight_layout()
diagram_path = '/Users/darvis01/Documents/ClonalGE/hdi_rope_diagram.png'
plt.savefig(diagram_path, dpi=200, bbox_inches='tight', facecolor='white')
plt.close()

# ── Build PowerPoint ──────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

GREEN = RGBColor(0x1a, 0x6b, 0x3a)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x55, 0x55, 0x55)
MAROON = RGBColor(0x6b, 0x1a, 0x1a)
BLUE = RGBColor(0x1a, 0x3a, 0x6b)


def add_text(slide, left, top, width, height, text, size=14, bold=False,
             color=DARK, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return tf


def add_multiline(slide, left, top, width, height, title, lines, title_size=14,
                  title_color=MAROON, line_size=12, line_color=DARK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(title_size)
    p.font.bold = True
    p.font.color.rgb = title_color
    for line in lines:
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(line_size)
        p.font.color.rgb = line_color
        p.space_before = Pt(4)


# ── Title ─────────────────────────────────────────────────────────────────────
add_text(slide, Inches(0.4), Inches(0.15), Inches(12), Inches(0.55),
         "Bayesian Differential Expression: HDI + ROPE", size=26, bold=True, color=BLUE)
add_text(slide, Inches(0.4), Inches(0.6), Inches(12), Inches(0.35),
         "Is gene g differentially expressed between clone k₁ and clone k₂?",
         size=15, color=GREY)

# ── Definitions box (top-right) ───────────────────────────────────────────────
add_multiline(slide, Inches(6.8), Inches(1.0), Inches(6.2), Inches(2.0),
              "Definitions", [
                  "",
                  "HDI (Highest Density Interval):",
                  "  The shortest interval containing X% of the posterior",
                  "  samples. Every point inside has higher density than",
                  "  any point outside. (We use 95%.)",
                  "",
                  "ROPE (Region of Practical Equivalence):",
                  "  An interval [−ε, +ε] around zero representing",
                  "  differences too small to be biologically meaningful.",
                  "  If the true difference is inside the ROPE, the two",
                  "  clones are considered practically equivalent.",
              ], title_size=15, title_color=BLUE, line_size=11)

# ── Steps (left column) ──────────────────────────────────────────────────────
y = Inches(1.05)
add_text(slide, Inches(0.4), y, Inches(6), Inches(0.3),
         "Step 1: Posterior samples of the difference", size=14, bold=True, color=GREEN)
add_text(slide, Inches(0.5), y + Inches(0.3), Inches(6), Inches(0.55),
         "From MCMC we have S samples. Subtract sample-by-sample:\n"
         "δ_g⁽ⁱ⁾ = B_{k₁,g}⁽ⁱ⁾ − B_{k₂,g}⁽ⁱ⁾,   i = 1, …, S",
         size=12)

y = Inches(2.0)
add_text(slide, Inches(0.4), y, Inches(6), Inches(0.3),
         "Step 2: Compute the 95% HDI", size=14, bold=True, color=GREEN)
add_text(slide, Inches(0.5), y + Inches(0.3), Inches(6), Inches(0.55),
         "Find the shortest interval [low, high] containing 95% of δ_g samples.\n"
         "→ Tells us where the true difference credibly lies.",
         size=12)

y = Inches(2.9)
add_text(slide, Inches(0.4), y, Inches(6), Inches(0.3),
         "Step 3: Define the ROPE", size=14, bold=True, color=GREEN)
add_text(slide, Inches(0.5), y + Inches(0.3), Inches(6), Inches(0.75),
         "ROPE = [−ε, +ε]  — the 'negligible difference' zone.\n"
         "ε = 10th percentile of |B_{k₁,g} − B_{k₂,g}| across all genes & pairs.\n"
         "→ Adaptive to data scale; bottom 10% of differences = negligible.",
         size=12)

y = Inches(4.0)
add_text(slide, Inches(0.4), y, Inches(6), Inches(0.3),
         "Step 4: Decision rule", size=14, bold=True, color=GREEN)
add_text(slide, Inches(0.5), y + Inches(0.3), Inches(6.2), Inches(0.85),
         "Gene is DE  ⟺  HDI_low > +ε   OR   HDI_high < −ε\n\n"
         "✅ Entire HDI outside ROPE → DE (credible AND meaningful)\n"
         "❌ HDI overlaps ROPE → Not DE (difference may be negligible)",
         size=12)

# ── Why box (bottom-right) ────────────────────────────────────────────────────
add_multiline(slide, Inches(6.8), Inches(3.5), Inches(6.2), Inches(1.8),
              "Why HDI + ROPE?", [
                  "• Purely Bayesian — no p-values, no frequentist tests",
                  "• Two requirements simultaneously:",
                  "    – Difference is credibly non-zero (HDI excludes 0)",
                  "    – Difference is large enough to matter (HDI outside ROPE)",
                  "• Avoids calling tiny but 'statistically certain' differences as DE",
              ], title_size=14, line_size=12)

# ── Diagram ───────────────────────────────────────────────────────────────────
slide.shapes.add_picture(diagram_path, Inches(2.8), Inches(5.15), width=Inches(7.8))

# ── Reference ─────────────────────────────────────────────────────────────────
add_text(slide, Inches(0.4), Inches(7.15), Inches(10), Inches(0.3),
         "Kruschke (2018), Rejecting or Accepting Parameter Values in Bayesian Estimation, AMPPS 1(2)",
         size=9, color=RGBColor(0x88, 0x88, 0x88))

# ── Save ──────────────────────────────────────────────────────────────────────
out_path = '/Users/darvis01/Documents/ClonalGE/HDI_ROPE_method.pptx'
prs.save(out_path)
print(f"Saved → {out_path}")
os.remove(diagram_path)
